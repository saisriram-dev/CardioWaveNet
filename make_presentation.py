"""
make_presentation.py - builds the 11-slide, light-themed review deck.

Run it LAST (after process_ecg.py, train.py and evaluate.py) so the result
slides contain your real curves, confusion matrix and metrics.
Output: outputs/CardioWaveNet_Review.pptx  (speaker notes on every slide)
"""
import os
import json

import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from config import CLASSES, IMAGE_DIR, OUTPUT_DIR

PRESENTER = "P Sai Sri Ram"  # edit if needed
DECK_PATH = os.path.join(OUTPUT_DIR, "CardioWaveNet_Review.pptx")
TOTAL_SLIDES = 11

# ---------------- light theme ----------------
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x6B, 0x72, 0x80)
RED = RGBColor(0xC0, 0x39, 0x2B)
RED_LIGHT = RGBColor(0xFB, 0xEC, 0xEA)
TEAL = RGBColor(0x0F, 0x76, 0x6E)
TEAL_LIGHT = RGBColor(0xE6, 0xF4, 0xF2)
PANEL = RGBColor(0xF4, 0xF6, 0xF9)
LINE = RGBColor(0xD5, 0xDB, 0xE3)
FONT = "Calibri"

CLASS_NAMES = {"N": "Normal (incl. bundle branch block)", "S": "Supraventricular ectopic",
               "V": "Ventricular ectopic", "F": "Fusion of normal + ventricular",
               "Q": "Paced / unclassifiable"}
CLASS_SHORT = {"N": "N · Normal", "S": "S · Supraventricular", "V": "V · Ventricular",
               "F": "F · Fusion", "Q": "Q · Paced/unknown"}
CLASS_SYMBOLS = {"N": "N, L, R, e, j", "S": "A, a, J, S", "V": "V, E", "F": "F", "Q": "/, f, Q"}

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


# ---------------- drawing helpers ----------------
def style_run(run, size, color=INK, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold


def text_box(slide, x, y, w, h, lines, size=16, color=INK, align=PP_ALIGN.LEFT,
             bullet=False, space=6):
    """lines: list of str or (bold_head, normal_body) tuples."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        head, body = line if isinstance(line, tuple) else ("", line)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        if bullet:
            r = p.add_run()
            r.text = "\u25AA  "
            style_run(r, size, RED, bold=True)
        if head:
            r = p.add_run()
            r.text = head
            style_run(r, size, color, bold=True)
        if body:
            r = p.add_run()
            r.text = body
            style_run(r, size, color)
    return tb


def box(slide, x, y, w, h, text="", fill=PANEL, color=INK, size=14, bold=False,
        border=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, align=PP_ALIGN.CENTER):
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = 0.08
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if border is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = border
        shp.line.width = Pt(1.25)
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    for i, line in enumerate(text.split("\n") if text else []):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        style_run(r, size, color, bold=(bold and i == 0))
    return shp


def arrow(slide, x, y):
    box(slide, x, y, 0.26, 0.3, fill=RED, shape=MSO_SHAPE.RIGHT_ARROW)


def panel(slide, x, y, w, h, heading, lines, head_color=TEAL, size=14, fill=PANEL):
    box(slide, x, y, w, h, fill=fill)
    text_box(slide, x + 0.2, y + 0.12, w - 0.4, 0.45, [(heading, "")], size=size + 3,
             color=head_color)
    text_box(slide, x + 0.2, y + 0.62, w - 0.4, h - 0.75, lines, size=size, bullet=True)


def table(slide, x, y, rows, col_w, row_h=0.45, size=12, highlight_row=None):
    n_rows, n_cols = len(rows), len(rows[0])
    shape = slide.shapes.add_table(n_rows, n_cols, Inches(x), Inches(y),
                                   Inches(sum(col_w)), Inches(row_h * n_rows))
    tbl = shape.table
    for j, w in enumerate(col_w):
        tbl.columns[j].width = Inches(w)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, value in enumerate(row):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            if i == 0:
                cell.fill.fore_color.rgb = TEAL
            elif i == highlight_row:
                cell.fill.fore_color.rgb = RED_LIGHT
            else:
                cell.fill.fore_color.rgb = PANEL if i % 2 else WHITE
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.text_frame.word_wrap = True
            run = cell.text_frame.paragraphs[0].add_run()
            run.text = str(value)
            style_run(run, size, WHITE if i == 0 else INK,
                      bold=(i == 0 or i == highlight_row))
    return shape


def picture(slide, path, x, y, max_w, max_h, missing_msg):
    if not path or not os.path.exists(path):
        box(slide, x, y, max_w, max_h, missing_msg, fill=PANEL, color=MUTED, size=14,
            border=LINE)
        return
    w, h = Image.open(path).size
    scale = min(max_w / w, max_h / h)
    pw, ph = w * scale, h * scale
    slide.shapes.add_picture(path, Inches(x + (max_w - pw) / 2),
                             Inches(y + (max_h - ph) / 2), Inches(pw), Inches(ph))


def metric_card(slide, x, y, w, h, value, label, color):
    shp = box(slide, x, y, w, h, fill=PANEL)
    tf = shp.text_frame
    tf.margin_left = Inches(0.25)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = value
    style_run(r, 28, color, bold=True)
    r = p.add_run()
    r.text = "   " + label
    style_run(r, 15, MUTED)


def ecg_trace(slide, x0, x1, y0, amp=0.8, beat_w=2.1):
    pattern = [(0.00, 0), (0.10, 0), (0.13, -0.10), (0.17, 0), (0.24, 0), (0.26, 0.08),
               (0.29, -1.0), (0.32, 0.30), (0.35, 0), (0.46, 0), (0.53, -0.20),
               (0.60, 0), (1.00, 0)]
    pts, x = [], x0
    while x + beat_w <= x1 + 1e-6:
        pts += [(x + fx * beat_w, y0 + fy * amp) for fx, fy in pattern]
        x += beat_w
    pts.append((x1, y0))
    fb = slide.shapes.build_freeform(Inches(pts[0][0]), Inches(pts[0][1]), scale=1.0)
    fb.add_line_segments([(Inches(px), Inches(py)) for px, py in pts[1:]], close=False)
    shp = fb.convert_to_shape()
    shp.fill.background()
    shp.line.color.rgb = RED
    shp.line.width = Pt(2.25)
    shp.shadow.inherit = False


def set_bg(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE


def footer(slide, num):
    text_box(slide, 0.5, 7.0, 7, 0.35, ["Cardio Wave Net  ·  ECG beat classification"],
             size=11, color=MUTED)
    text_box(slide, 10.85, 7.0, 2.0, 0.35, [f"{num} / {TOTAL_SLIDES}"], size=11,
             color=MUTED, align=PP_ALIGN.RIGHT)


def new_slide(title, num, subtitle=None):
    s = prs.slides.add_slide(BLANK)
    set_bg(s)
    box(s, 0, 0, 0.12, 7.5, fill=RED, shape=MSO_SHAPE.RECTANGLE)
    text_box(s, 0.5, 0.3, 12.3, 0.75, [(title, "")], size=30)
    box(s, 0.55, 1.02, 1.1, 0.06, fill=TEAL, shape=MSO_SHAPE.RECTANGLE)
    if subtitle:
        text_box(s, 0.5, 1.1, 12.3, 0.45, [subtitle], size=14, color=MUTED)
    footer(s, num)
    return s


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


# ---------------- load results ----------------
def load_counts():
    counts = {}
    for c in CLASSES:
        folder = os.path.join(IMAGE_DIR, c)
        counts[c] = (len([f for f in os.listdir(folder) if f.endswith(".png")])
                     if os.path.isdir(folder) else 0)
    return counts


def load_history():
    path = os.path.join(OUTPUT_DIR, "history.json")
    if not os.path.exists(path):
        return None
    with open(path) as f:
        h = json.load(f)
    best = int(np.argmin(h["val_loss"]))
    return {"epochs": len(h["loss"]), "best_epoch": best + 1,
            "val_loss": h["val_loss"][best], "val_acc": h["val_accuracy"][best],
            "train_acc": h["accuracy"][best]}


def load_report():
    path = os.path.join(OUTPUT_DIR, "classification_report.txt")
    rows, summary = [], {}
    if not os.path.exists(path):
        return rows, summary
    with open(path) as f:
        for line in f:
            p = line.split()
            if len(p) == 5 and p[0] in CLASSES:
                rows.append(p)
            elif p and p[0] == "accuracy":
                summary["accuracy"] = float(p[1])
            elif len(p) == 6 and p[1] == "avg":
                summary[p[0]] = float(p[4])
    return rows, summary


def load_cm():
    path = os.path.join(OUTPUT_DIR, "confusion_matrix.csv")
    return np.loadtxt(path, delimiter=",", dtype=int) if os.path.exists(path) else None


def sample_image(cls):
    folder = os.path.join(IMAGE_DIR, cls)
    if not os.path.isdir(folder):
        return None
    files = sorted(f for f in os.listdir(folder) if f.endswith(".png"))
    return os.path.join(folder, files[len(files) // 2]) if files else None


def pct(x):
    return f"{x * 100:.1f}%"


counts = load_counts()
n_total = sum(counts.values())
hist = load_history()
report_rows, summary = load_report()
cm = load_cm()

n_train_txt = f"~{int(0.7 * n_total):,}" if n_total else "~20,000"
if hist:
    gap = hist["train_acc"] - hist["val_acc"]
    train_txt = (f"Stopped after {hist['epochs']} epochs; best epoch {hist['best_epoch']} "
                 f"(val accuracy {pct(hist['val_acc'])}, val loss {hist['val_loss']:.3f}).")
    gap_txt = (f"Small train-val gap ({gap * 100:.1f} pp) at the best epoch"
               if gap < 0.03 else
               f"Train-val gap of {gap * 100:.1f} pp, controlled by early stopping")
else:
    train_txt = "Run train.py to fill in the training summary."
    gap_txt = "Learning curves monitored for overfitting (slide 9)"

if cm is not None:
    off = cm.copy()
    np.fill_diagonal(off, 0)
    i, j = np.unravel_index(off.argmax(), off.shape)
    confusion_txt = (f"Most frequent error: true {CLASSES[i]} predicted as {CLASSES[j]} "
                     f"({off[i, j]} test beats).")
else:
    confusion_txt = "Run evaluate.py to fill in the confusion analysis."

acc_txt = pct(summary["accuracy"]) if "accuracy" in summary else "—"
macro_txt = f"{summary['macro']:.3f}" if "macro" in summary else "—"
weighted_txt = f"{summary['weighted']:.3f}" if "weighted" in summary else "—"


# ================= Slide 1: Introduction =================
s = prs.slides.add_slide(BLANK)
set_bg(s)
box(s, 0, 0, 0.12, 7.5, fill=RED, shape=MSO_SHAPE.RECTANGLE)
text_box(s, 0.6, 0.4, 12.2, 1.0, [("Cardio Wave Net", "")], size=54)
text_box(s, 0.62, 1.35, 12.2, 0.5,
         ["Heartbeat classification from ECG wavelet scalograms with a compact 2-D CNN"],
         size=20, color=TEAL)
ecg_trace(s, 0.6, 12.75, y0=2.95, amp=0.8)
intro_cols = [
    ("The problem", RED, [
        "Arrhythmia diagnosis needs beat-by-beat review of long Holter ECGs (~100,000 beats/day).",
        "Manual review is slow and tiring; rare abnormal beats are easy to miss.",
    ]),
    ("Our approach", TEAL, [
        "Cut each beat around its R-peak and convert it to a CWT scalogram (time x frequency image).",
        "Classify it with CardioWaveNet into the 5 AAMI EC57 classes: N, S, V, F, Q.",
    ]),
    ("Since mid-review", TEAL, [
        "Scaled from 4 records to all 48 MIT-BIH records.",
        "Designed, trained and evaluated CardioWaveNet; built a Streamlit demo.",
    ]),
]
for k, (head, color, lines) in enumerate(intro_cols):
    panel(s, 0.6 + k * 4.12, 3.55, 3.95, 2.8, head, lines, head_color=color, size=13)
text_box(s, 0.6, 6.5, 8, 0.4, [("Presented by: ", PRESENTER)], size=14)
footer(s, 1)
notes(s, "Arrhythmias are abnormal heartbeats. A 24-hour Holter recording contains around "
         "100,000 beats, so manual review is slow. CardioWaveNet automates beat-level "
         "classification: each beat becomes a CWT scalogram, and a compact 2-D CNN assigns "
         "one of the 5 AAMI classes. At mid-review I used 4 records; for this review I scaled "
         "to all 48 records, trained and evaluated the model, and built a live demo.")

# ================= Slide 2: Literature review =================
s = new_slide("Literature Review", 2,
              "From hand-crafted features to CNNs on beat images and scalograms")
lit_rows = [
    ["Study", "Method", "Relevance to CardioWaveNet"],
    ["de Chazal et al., 2004 (IEEE TBME)",
     "Hand-crafted morphology + RR-interval features, linear discriminant classifier",
     "Set up the inter-patient (DS1/DS2) protocol; features need expert design"],
    ["Kiranyaz et al., 2016 (IEEE TBME)", "Compact 1-D CNN, patient-specific training",
     "Small CNNs learn features directly from beats; cheap enough for real time"],
    ["Acharya et al., 2017 (Comput. Biol. Med.)",
     "9-layer 1-D CNN, 5 AAMI classes, synthetic data to balance classes",
     "~94% accuracy; shows class imbalance must be handled explicitly"],
    ["Kachuee et al., 2018 (IEEE ICHI)", "Deep residual 1-D CNN on MIT-BIH, 5 AAMI classes",
     "~93% average accuracy; learned features transfer to MI detection"],
    ["Jun et al., 2018 (arXiv)",
     "Each beat drawn as a 128x128 grayscale image, VGG-style 2-D CNN",
     "2-D CNNs on beat images work well; motivates our 128x128 input"],
    ["Wang et al., 2021 (Entropy)", "CWT scalograms + 2-D CNN, with RR-interval features",
     "Closest to our work: scalograms expose time-frequency patterns to a CNN"],
    ["This work: CardioWaveNet",
     "CWT (Morlet) scalogram, 4-block 2-D CNN with GAP, ~107k parameters",
     "All 48 records, AAMI 5-class, class-weighted loss, live Streamlit demo"],
]
table(s, 0.5, 1.65, lit_rows, [3.0, 4.5, 4.8], row_h=0.6, size=12,
      highlight_row=len(lit_rows) - 1)
notes(s, "The field moved from hand-crafted features (de Chazal) to 1-D CNNs (Kiranyaz, "
         "Acharya, Kachuee) to 2-D CNNs on beat images and scalograms (Jun, Wang). We follow "
         "the 2-D scalogram line but keep the network compact.\n\nReferences:\n"
         "de Chazal P, O'Dwyer M, Reilly RB. IEEE TBME 51(7):1196-1206, 2004.\n"
         "Kiranyaz S, Ince T, Gabbouj M. IEEE TBME 63(3):664-675, 2016.\n"
         "Acharya UR et al. Computers in Biology and Medicine 89:389-396, 2017.\n"
         "Kachuee M, Fazeli S, Sarrafzadeh M. IEEE ICHI 2018.\n"
         "Jun TJ et al. arXiv:1804.06812, 2018.\n"
         "Wang T et al. Entropy 23(1):119, 2021.\n"
         "Moody GB, Mark RG. IEEE EMB Magazine 20(3):45-50, 2001 (MIT-BIH).")

# ================= Slide 3: Dataset =================
s = new_slide("Dataset & Beat Classes", 3,
              "MIT-BIH Arrhythmia Database mapped to the AAMI EC57 standard")
ds_rows = [["Class", "Meaning", "MIT-BIH labels", "Images used"]]
for c in CLASSES:
    ds_rows.append([c, CLASS_NAMES[c], CLASS_SYMBOLS[c], f"{counts[c]:,}" if counts[c] else "—"])
ds_rows.append(["Total", "", "", f"{n_total:,}" if n_total else "—"])
table(s, 0.5, 1.65, ds_rows, [0.9, 2.9, 2.1, 1.7], row_h=0.5, size=13,
      highlight_row=len(ds_rows) - 1)
text_box(s, 0.5, 5.35, 7.6, 1.4, [
    "Normal beats are kept 1 in 10 at extraction (N is ~80% of all MIT-BIH beats).",
    "Remaining imbalance is handled with a class-weighted loss (slide 8).",
], size=14, bullet=True)
panel(s, 8.4, 1.65, 4.45, 5.0, "MIT-BIH at a glance", [
    "48 half-hour, two-channel ambulatory ECG recordings from 47 subjects",
    "360 Hz sampling, 11-bit resolution",
    "~110,000 beats annotated by two or more cardiologists",
    "Channel 0 used (modified lead II in most records)",
    "Mid-review: 4 records. Now: all 48 records",
], size=14)
notes(s, "AAMI EC57 groups the MIT-BIH beat labels into 5 classes, which is the standard "
         "way to report results so they are comparable across papers. The image counts in "
         "the table come straight from the generated scalogram folders.")

# ================= Slide 4: Pipeline =================
s = new_slide("Methodology: ECG Beat to Scalogram", 4,
              "Every beat becomes a 128x128 time-frequency image")
steps = ["MIT-BIH ECG\n48 records · 360 Hz", "Beat window\n-100 / +150 samples\naround R-peak",
         "CWT\nMorlet wavelet\nscales 1-63", "Scalogram\n|coef| scaled 0-255\n128x128 PNG",
         "CardioWaveNet\n2-D CNN", "5 AAMI classes\nN · S · V · F · Q"]
x = 0.6
for k, txt in enumerate(steps):
    fill, border = (RED_LIGHT, RED) if k == 4 else (TEAL_LIGHT, TEAL)
    box(s, x, 1.7, 1.7, 1.15, txt, fill=fill, border=border, size=12, bold=True)
    if k < len(steps) - 1:
        arrow(s, x + 1.76, 2.12)
    x += 2.08
panel(s, 0.5, 3.15, 6.6, 3.65, "Why a CWT scalogram?", [
    ("FFT: ", "shows which frequencies exist but loses when they occur."),
    ("STFT: ", "one fixed window cannot resolve both the sharp QRS and slower waves."),
    ("CWT: ", "multi-resolution. Short wavelets (small scales) resolve the sharp QRS; long "
              "wavelets (large scales) capture slower components."),
    ("Morlet: ", "a smooth oscillating wavelet that resembles QRS morphology; widely used "
                 "for ECG scalograms."),
], size=14)
box(s, 7.3, 3.15, 5.55, 3.65, fill=PANEL)
text_box(s, 7.5, 3.25, 5.2, 0.45, [("Example scalograms", "")], size=17, color=TEAL)
for k, c in enumerate(["N", "S", "V"]):
    picture(s, sample_image(c), 7.5 + k * 1.75, 3.75, 1.6, 1.6, f"{c}\n(run\nprocess_ecg.py)")
    text_box(s, 7.5 + k * 1.75, 5.38, 1.6, 0.35, [CLASS_SHORT[c].split(" · ")[1]], size=12,
             color=MUTED, align=PP_ALIGN.CENTER)
text_box(s, 7.5, 5.8, 5.2, 0.95, [
    "Rows = scale, columns = time. Window: 100 samples (278 ms) before the R-peak keeps the "
    "P wave; 150 samples (417 ms) after keeps the ST segment and T wave."], size=12)
notes(s, "Each beat is cut 100 samples before and 150 after the annotated R-peak, about "
         "0.7 s in total, which covers P wave, QRS and T wave at normal heart rates. The CWT "
         "with a Morlet wavelet gives a time-frequency image: sharp, high-frequency QRS energy "
         "shows at small scales, slower content at larger scales. Magnitudes are min-max scaled "
         "to 0-255 and resized to 128x128.")

# ================= Slide 5: Architecture =================
s = new_slide("CardioWaveNet Architecture", 5,
              "4 convolutional blocks + global average pooling + small classifier head")
blocks = [
    ("Input\n128×128×1\n÷255 rescale", 2.9, PANEL, MUTED),
    ("Block 1\nConv 3×3, 16\nReLU → BN\nMaxPool 2×2\n64×64×16", 2.6, TEAL_LIGHT, TEAL),
    ("Block 2\nConv 3×3, 32\nReLU → BN\nMaxPool 2×2\n32×32×32", 2.3, TEAL_LIGHT, TEAL),
    ("Block 3\nConv 3×3, 64\nReLU → BN\nMaxPool 2×2\n16×16×64", 2.0, TEAL_LIGHT, TEAL),
    ("Block 4\nConv 3×3, 128\nReLU → BN\n(no pool)\n16×16×128", 2.0, TEAL_LIGHT, TEAL),
    ("Global Avg\nPool\n128", 1.4, RED_LIGHT, RED),
    ("Dropout 0.4\nDense 64\nReLU", 1.4, RED_LIGHT, RED),
    ("Dense 5\nSoftmax\nN S V F Q", 1.2, RED_LIGHT, RED),
]
x, cy, bw = 0.55, 3.15, 1.2
for k, (txt, h, fill, border) in enumerate(blocks):
    box(s, x, cy - h / 2, bw, h, txt, fill=fill, border=border, size=10.5, bold=True)
    if k < len(blocks) - 1:
        arrow(s, x + bw + 0.06, cy - 0.15)
    x += bw + 0.38
text_box(s, 2.13, 4.72, 4.94, 0.4, [("Feature extractor (4 conv blocks)", "")], size=14,
         color=TEAL, align=PP_ALIGN.CENTER)
text_box(s, 8.45, 4.72, 4.36, 0.4, [("Classifier head", "")], size=14, color=RED,
         align=PP_ALIGN.CENTER)
cards = [
    ("106,693 parameters", "≈0.43 MB of weights: trains on a laptop CPU, real-time inference"),
    ("Receptive field 38×38 px", "≈206 ms of the beat: wide enough for a wide ventricular QRS"),
    ("GAP instead of Flatten", "Flatten + Dense(64) alone would need 2.1 M weights (≈20× the whole model)"),
]
for k, (head, body) in enumerate(cards):
    box(s, 0.5 + k * 4.15, 5.25, 4.0, 1.5, fill=PANEL)
    text_box(s, 0.7 + k * 4.15, 5.35, 3.6, 1.35, [(head, ""), body], size=14)
notes(s, "Input scalograms are rescaled to 0-1. Each of the first three blocks is a 3x3 "
         "convolution, ReLU, batch normalisation and 2x2 max pooling, halving the resolution "
         "while doubling the filters. Block 4 has no pooling, keeping a 16x16 map. Global "
         "average pooling turns 16x16x128 into 128 numbers, followed by dropout 0.4, a "
         "64-unit dense layer and a 5-way softmax. Total 106,693 parameters; this can be "
         "verified with `python model.py`.")

# ================= Slide 6: Why (1) =================
s = new_slide("Why This Architecture? (1) Feature Extractor", 6,
              "Design derived from the input image and ECG physiology")
text_box(s, 0.5, 1.65, 5.6, 4.5, [
    ("Why a 2-D CNN? ", "A scalogram is an image whose diagnostic cues are local shapes: the "
                        "QRS energy blob, its width and frequency spread. Convolutions share "
                        "weights and tolerate small R-peak misalignments."),
    ("Why 3×3 kernels? ", "Two stacked 3×3 layers see a 5×5 region with 18 instead of 25 "
                          "weights per channel pair, plus an extra non-linearity (VGG principle)."),
    ("Why 16→32→64→128? ", "Each pooling halves resolution, so filters double to keep capacity "
                           "balanced: edges early, whole-beat patterns late."),
    ("Why 4 blocks? ", "Chosen from the receptive field (table)."),
], size=14, bullet=True, space=10)
text_box(s, 6.4, 1.6, 6.45, 0.45, [("Receptive field vs. ECG physiology", "")], size=17,
         color=TEAL)
rf_rows = [
    ["Depth", "RF (px)", "≈ Time", "What fits inside"],
    ["Block 1", "3", "16 ms", "Slopes, sharp edges"],
    ["Block 2", "8", "43 ms", "Wave fragments (R upstroke)"],
    ["Block 3", "18", "98 ms", "A normal QRS (80-100 ms)"],
    ["Block 4 (used)", "38", "206 ms", "Wide ventricular QRS (>120 ms) + context"],
    ["Block 5 (not used)", "78", "424 ms", "Most of the beat; +0.3 M params (≈4× model)"],
]
table(s, 6.4, 2.1, rf_rows, [1.6, 0.9, 0.9, 3.05], row_h=0.55, size=12, highlight_row=4)
text_box(s, 6.4, 5.5, 6.45, 0.6, [
    "1 px on the time axis = 250 samples ÷ 128 px ÷ 360 Hz ≈ 5.4 ms."], size=12, color=MUTED)
box(s, 0.5, 6.15, 12.35, 0.7,
    "4 blocks = the shallowest network whose receptive field covers a wide ventricular QRS; "
    "a 5th block costs ~4× the parameters for context the task does not need.",
    fill=RED_LIGHT, size=14, bold=True)
notes(s, "Receptive field is computed with r_out = r_in + (k-1)*j_in and j_out = j_in*s. "
         "After block 3 a unit sees about 98 ms, roughly a normal QRS. Ventricular beats have "
         "wide QRS complexes above 120 ms, so we need block 4, which sees about 206 ms. A "
         "fifth block would see 424 ms but almost quadruple the parameters. These are "
         "theoretical receptive fields; the effective field is smaller and centred, which is "
         "another reason not to go shallower than 4 blocks.")

# ================= Slide 7: Why (2) =================
s = new_slide("Why This Architecture? (2) Head & Regularisation", 7,
              "Each choice compared with the obvious alternative")
why_rows = [
    ["Design choice", "Alternative", "Why we chose it"],
    ["Global Average Pooling", "Flatten + Dense",
     f"Flatten gives 32,768 features → 2.1 M weights in one layer for {n_train_txt} training "
     "images. GAP gives 128 features (8,256 weights): less overfitting, position-tolerant."],
    ["BatchNorm after every conv", "No normalisation",
     "Keeps activations well-scaled so Adam can use lr = 1e-3, converges in fewer epochs, "
     "adds mild regularisation."],
    ["Dropout 0.4, head only", "None / everywhere",
     "Overfitting risk concentrates in the dense head; conv layers are already regularised "
     "by BN and weight sharing."],
    ["Compact custom CNN (~0.1 M)", "ImageNet ResNet-50 (~25 M) / VGG-16 (~138 M)",
     "Scalograms are 1-channel, non-natural images; modest dataset; must train on a CPU and "
     "run live in Streamlit."],
    ["Softmax + sparse cross-entropy", "Sigmoid / one-vs-rest",
     "AAMI classes are mutually exclusive; labels are stored as integers 0-4."],
]
table(s, 0.5, 1.65, why_rows, [2.4, 2.8, 7.15], row_h=0.62, size=12)
box(s, 0.5, 5.6, 12.35, 1.25, fill=TEAL_LIGHT, border=TEAL)
text_box(s, 0.7, 5.65, 12.0, 1.2, [
    ("Evidence for the design (formal ablation planned for the final review)", ""),
    f"① Receptive field sized to QRS physiology   ② ~107k parameters suits {n_train_txt} "
    f"training images   ③ {gap_txt}   ④ Follows proven CNN-for-ECG designs "
    "(Kiranyaz 2016, Jun 2018, Wang 2021)",
], size=13, color=INK)
notes(s, "I have not run an ablation study yet, so I justify each choice by reasoning and "
         "evidence. GAP versus Flatten is the biggest one: Flatten would put 2.1 million "
         "weights in a single layer, about 20 times the whole model. Batch norm stabilises "
         "training. Dropout sits only in the head. Transfer learning from ImageNet was not "
         "used because scalograms are single-channel and look nothing like natural photos, "
         "and those models are 200 to 1000 times larger. The planned ablation is listed on "
         "the final slide.")

# ================= Slide 8: Training =================
s = new_slide("Training Strategy", 8, "Handling class imbalance and avoiding overfitting")
text_box(s, 0.5, 1.65, 6.7, 5.1, [
    ("Split: ", "stratified 70 / 15 / 15 (train / val / test), fixed seed 42; the test set is "
                "never seen during training."),
    ("Imbalance, step 1: ", "keep 1 in 10 normal beats at extraction."),
    ("Imbalance, step 2: ", "'balanced' class weights w_c = n / (K · n_c): rare classes cost "
                            "more when misclassified, without duplicating images."),
    ("Optimiser: ", "Adam, lr 1e-3, batch 64, up to 30 epochs."),
    ("Callbacks: ", "EarlyStopping on val loss (patience 5, restore best), ReduceLROnPlateau "
                    "(×0.5 after 2 flat epochs), ModelCheckpoint keeps the best model."),
    ("This run: ", train_txt),
], size=14, bullet=True, space=10)
text_box(s, 7.5, 1.6, 5.35, 0.45, [("Class weights (≈, from image counts)", "")], size=17,
         color=TEAL)
w_rows = [["Class", "Images", "Weight"]]
for c in CLASSES:
    w = n_total / (len(CLASSES) * counts[c]) if counts[c] else None
    w_rows.append([CLASS_SHORT[c], f"{counts[c]:,}" if counts[c] else "—",
                   f"{w:.2f}" if w else "—"])
table(s, 7.5, 2.1, w_rows, [2.35, 1.5, 1.5], row_h=0.48, size=13)
box(s, 7.5, 5.15, 5.35, 1.6,
    "Why not oversampling / SMOTE?\nDuplicated or synthetic scalograms risk memorisation; "
    "weighting the loss is free and leaves the data untouched.",
    fill=PANEL, size=13, bold=True)
notes(s, "Class weights are computed by scikit-learn's 'balanced' rule. Early stopping "
         "restores the epoch with the lowest validation loss, and the learning rate is "
         "halved whenever validation loss stalls for two epochs.")

# ================= Slide 9: Results 1 =================
s = new_slide("Results: Learning Curves & Confusion Matrix", 9,
              "Held-out test set (15%), never used for training or model selection")
picture(s, os.path.join(OUTPUT_DIR, "training_curves.png"), 0.5, 1.65, 7.2, 2.9,
        "training_curves.png missing - run train.py")
panel(s, 0.5, 4.75, 7.2, 2.05, "What the curves show", [
    train_txt,
    gap_txt + ".",
    confusion_txt,
], size=13)
picture(s, os.path.join(OUTPUT_DIR, "confusion_matrix.png"), 7.95, 1.6, 4.9, 5.2,
        "confusion_matrix.png missing - run evaluate.py")
notes(s, "Left: accuracy and loss for training and validation. Training accuracy can sit "
         "below validation accuracy because dropout is active during training. Right: the "
         "test confusion matrix; the diagonal is correct predictions. S beats look like N "
         "beats in shape; what separates them is prematurity (RR interval), which a "
         "single-beat image cannot see, so S-N confusion is expected.")

# ================= Slide 10: Results 2 =================
s = new_slide("Results: Per-class Test Metrics & Demo", 10,
              "Macro-F1 is the headline metric for imbalanced classes")
if report_rows:
    m_rows = [["Class", "Precision", "Recall", "F1", "Support"]]
    m_rows += [[CLASS_SHORT[p[0]], p[1], p[2], p[3], p[4]] for p in report_rows]
    table(s, 0.5, 1.65, m_rows, [2.6, 1.0, 1.0, 1.0, 1.1], row_h=0.5, size=13)
    best = max(report_rows, key=lambda p: float(p[3]))
    worst = min(report_rows, key=lambda p: float(p[3]))
    insight = [
        f"Strongest class: {best[0]} (F1 {best[3]}). Weakest: {worst[0]} "
        f"(F1 {worst[3]}, only {worst[4]} test beats).",
        "Macro-F1 weights every class equally; accuracy is dominated by the large classes.",
    ]
else:
    box(s, 0.5, 1.65, 6.7, 3.0, "classification_report.txt missing - run evaluate.py",
        fill=PANEL, color=MUTED, border=LINE)
    insight = ["Run evaluate.py to fill in per-class metrics."]
text_box(s, 0.5, 4.85, 6.7, 1.9, insight, size=14, bullet=True)
metric_card(s, 7.7, 1.65, 5.15, 0.95, acc_txt, "Test accuracy", TEAL)
metric_card(s, 7.7, 2.72, 5.15, 0.95, macro_txt, "Macro F1", RED)
metric_card(s, 7.7, 3.79, 5.15, 0.95, weighted_txt, "Weighted F1", TEAL)
panel(s, 7.7, 4.95, 5.15, 1.85, "Live demo: streamlit run app.py", [
    "Random beat of any class or an uploaded scalogram → predicted class, confidence and "
    "probability chart.",
], size=13, fill=TEAL_LIGHT)
notes(s, "Report macro-F1 first: because N, V and Q dominate the test set, plain accuracy can "
         "look high even if the rare F class is poor. Then switch to the Streamlit demo: "
         "pick a random V beat, show the prediction and probabilities, then open the "
         "Model performance tab.")

# ================= Slide 11: Conclusion =================
s = new_slide("Conclusion, Limitations & Future Work", 11)
f_count = f"{counts['F']:,}" if counts["F"] else "very few"
concl = [
    ("Conclusions", TEAL, [
        "Full pipeline on all 48 MIT-BIH records: R-peak window → CWT scalogram → CNN → AAMI class.",
        "CardioWaveNet: 4 conv blocks + GAP, ~107k parameters; depth set by a QRS-sized receptive field.",
        f"Held-out test: accuracy {acc_txt}, macro-F1 {macro_txt}.",
        "Interactive Streamlit demo for single-beat inference.",
    ]),
    ("Limitations", RED, [
        "Beat-level random split: beats of the same patient appear in train and test, so scores "
        "are optimistic vs. inter-patient testing.",
        "No ablation study yet; design justified by analysis and literature.",
        "Single beat, single lead: no RR-interval (rhythm) context, the main cue for S beats.",
        f"F class is tiny ({f_count} images).",
    ]),
    ("Future work", TEAL, [
        "Inter-patient evaluation with the de Chazal DS1/DS2 split.",
        "Ablation: 3 / 4 / 5 blocks, GAP vs Flatten, with/without BN and dropout.",
        "Add RR-interval features as a second input branch.",
        "Grad-CAM on scalograms to explain decisions; wider CWT scale range for slow P/T waves.",
    ]),
]
for k, (head, color, lines) in enumerate(concl):
    panel(s, 0.5 + k * 4.15, 1.4, 4.05, 4.95, head, lines, head_color=color, size=13)
text_box(s, 0.5, 6.42, 12.35, 0.5, [("Thank you  ·  Questions?", "")], size=20,
         color=TEAL, align=PP_ALIGN.CENTER)
notes(s, "Be upfront about the split: a random beat-level split lets neighbouring beats of "
         "the same patient land in both train and test, which inflates scores. Most papers "
         "reporting 98-99% use this setup; the stricter inter-patient DS1/DS2 split is the "
         "next step, together with the ablation study and RR-interval features.")

os.makedirs(OUTPUT_DIR, exist_ok=True)
prs.save(DECK_PATH)
print(f"Saved {DECK_PATH}")
